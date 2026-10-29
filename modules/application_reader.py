from io import BytesIO
from pathlib import Path
import re
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageEnhance, ImageFilter


def _cell_after(ws, label, value_col):
    for row in ws.iter_rows():
        for c in row:
            if label in str(c.value or ''):
                v = ws.cell(c.row, value_col).value
                return '' if v is None else str(v).strip()
    return ''


def read_xlsx(raw):
    ws = load_workbook(BytesIO(raw), data_only=True).worksheets[0]
    return {
        'customer': _cell_after(ws, '고객사', 8),
        'part_name': _cell_after(ws, '품목명', 26),
        'order_no': _cell_after(ws, '수주번호', 8),
        'po_step_no': _cell_after(ws, '발주/공정번호', 26),
        'source': 'Excel 신청서',
        'review_fields': [],
    }


def _render_pdf(raw):
    import fitz
    doc = fitz.open(stream=raw, filetype='pdf')
    page = doc[0]
    # Render near 300 dpi. The old version used a low-resolution page image.
    pix = page.get_pixmap(matrix=fitz.Matrix(4.0, 4.0), alpha=False)
    return Image.open(BytesIO(pix.tobytes('png'))).convert('RGB')


def _ocr_data(img, lang, psm=7):
    import pytesseract
    try:
        data = pytesseract.image_to_data(
            img, lang=lang, config=f'--oem 3 --psm {psm}',
            output_type=pytesseract.Output.DICT
        )
    except Exception:
        data = pytesseract.image_to_data(
            img, lang='eng', config=f'--oem 3 --psm {psm}',
            output_type=pytesseract.Output.DICT
        )
    words, confs = [], []
    for text, conf in zip(data.get('text', []), data.get('conf', [])):
        text = str(text).strip()
        try:
            score = float(conf)
        except Exception:
            score = -1
        if text and score >= 0:
            words.append(text)
            confs.append(score)
    return ' '.join(words).strip(), (sum(confs) / len(confs) if confs else 0)


def _prepare_crop(img, box):
    w, h = img.size
    crop = img.crop((int(box[0]*w), int(box[1]*h), int(box[2]*w), int(box[3]*h)))
    crop = ImageOps.grayscale(crop)
    crop = ImageEnhance.Contrast(crop).enhance(2.2)
    crop = crop.filter(ImageFilter.SHARPEN)
    # Remove table borders close to the crop edge and upscale.
    pad = max(6, int(min(crop.size) * 0.05))
    crop = crop.crop((pad, pad, max(pad+1, crop.width-pad), max(pad+1, crop.height-pad)))
    crop = ImageOps.expand(crop, border=24, fill='white')
    return crop.resize((crop.width*2, crop.height*2), Image.Resampling.LANCZOS)


def _clean_value(field, text):
    text = re.sub(r'\s+', ' ', text or '').strip(' |:_-.,')
    if field in ('order_no', 'po_step_no'):
        text = text.upper().replace(' ', '')
        text = re.sub(r'[^A-Z0-9,/-]', '', text)
        # A right-hand table border can be read as a trailing I in PS-number fields.
        if field == 'po_step_no' and re.fullmatch(r'PS\d{5,}I', text):
            text = text[:-1]
    return text


def _valid(field, value, confidence):
    if not value:
        return False
    if field == 'order_no':
        return confidence >= 35 and bool(re.fullmatch(r'[A-Z][A-Z0-9-]{4,20}', value))
    if field == 'po_step_no':
        return confidence >= 35 and bool(re.fullmatch(r'[A-Z0-9][A-Z0-9,/-]{4,40}', value))
    if field == 'part_name':
        return confidence >= 30 and len(value) >= 2
    if field == 'customer':
        return confidence >= 35 and len(value) >= 2
    return False


# Relative coordinates for both supplied old and new forms.
# Each field is read separately, so labels, signatures, and neighboring cells do not mix in.
FIELD_BOXES = {
    'order_no':  (0.225, 0.247, 0.498, 0.292),
    'customer':  (0.225, 0.286, 0.498, 0.326),
    'part_name': (0.708, 0.286, 0.985, 0.326),
    'po_step_no':(0.708, 0.320, 0.985, 0.359),
}


def _read_candidate(img):
    values, scores = {}, {}
    for field, box in FIELD_BOXES.items():
        crop = _prepare_crop(img, box)
        lang = 'eng' if field in ('order_no', 'po_step_no') else 'kor+eng'
        text, confidence = _ocr_data(crop, lang, psm=6 if field in ('order_no', 'po_step_no') else 7)
        values[field] = _clean_value(field, text)
        scores[field] = confidence
    orientation_score = 0
    if re.fullmatch(r'[A-Z][A-Z0-9-]{4,20}', values['order_no'] or ''):
        orientation_score += 5
    if re.fullmatch(r'[A-Z0-9][A-Z0-9-]{4,24}', values['po_step_no'] or ''):
        orientation_score += 4
    if values['part_name']:
        orientation_score += 2
    if values['customer']:
        orientation_score += 1
    orientation_score += sum(scores.values()) / 400
    return values, scores, orientation_score


def read_pdf(raw):
    original = _render_pdf(raw)
    candidates = []
    for deg in (0, 90, 180, 270):
        image = original if deg == 0 else original.rotate(deg, expand=True)
        values, scores, orientation_score = _read_candidate(image)
        candidates.append((orientation_score, deg, image, values, scores))
    _, rotation, corrected, values, scores = max(candidates, key=lambda x: x[0])

    result = {'source': 'PDF 신청서 항목별 OCR', 'rotation': rotation, 'review_fields': []}
    for field in ('customer', 'part_name', 'order_no', 'po_step_no'):
        value = values[field]
        if _valid(field, value, scores[field]):
            result[field] = value
        else:
            # Blank is safer than filling obvious OCR noise into the report.
            result[field] = ''
            result['review_fields'].append(field)
    return result, corrected


def read_application(uploaded):
    raw = uploaded.getvalue()
    ext = Path(uploaded.name).suffix.lower()
    if ext == '.xlsx':
        return read_xlsx(raw), None
    if ext == '.pdf':
        return read_pdf(raw)
    return {}, None
